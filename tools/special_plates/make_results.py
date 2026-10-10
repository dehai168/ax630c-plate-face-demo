exec(open('/workspace/special_plates/eval_special.py').read())
import csv,json,openpyxl,os,collections
from openpyxl.styles import Font,PatternFill,Alignment
os.chdir('/workspace/special_plates')
G=load_gt(); P=parse('runs/all_p3.tsv'); C=parse('runs/all_crop.tsv'); P2=parse('runs/gt_p2.tsv')
imgs=[os.path.basename(l.strip()) for l in open('list_all.txt')]
def infer_cat(t):
    if not t: return '未知(未检出)'
    if t.startswith('WJ'): return '武警(WJ)'
    if t.endswith('警'): return '警用(白)'
    if '使' in t or '领' in t or t[0].isdigit() or t.endswith('澳') or t.endswith('港'): return '使馆(黑,使)'
    if len(t)==8:
        if '警' in t: return '警用(白)'
        if not strict_ok(t): return '识别异常(8位不合法)'
        return '新能源小型(绿)' if t[2].isalpha() else '新能源大型(黄绿)'
    return '普通/其他'
import subprocess
def strict_ok(t):
    return bool(re.match(r'^[京津沪渝冀豫云辽黑湘皖鲁新苏浙赣鄂桂甘晋蒙陕吉闽贵粤青藏川宁琼][A-HJ-NP-Z](?:[DABCEFGHJK][A-HJ-NP-Z0-9][0-9]{4}|[A-HJ-NP-Z0-9]{5}[DABCEFGHJK])$',t))
rows=[]
for i,img in enumerate(imgs,1):
    r=P.get(img,{'read_fail':True,'best':None}); b=r['best']; cb=(C.get(img) or {}).get('best')
    gt,x=G.get(img,(None,False))
    t=b['text'] if b else ''
    note=[]
    if r['read_fail']: note.append('图片损坏无法读取')
    elif not b: note.append('未检出车牌')
    if gt and x: note.append('真值不确定(遮挡/模糊)，不计入统计')
    if b and not b.get('strict'): note.append('格式校验不通过(demo中不计票)')
    if gt and not x and b and t!=gt:
        if gt.endswith('警') and not t.endswith('警'): note.append('“警”字未识别')
        if '使' in gt: note.append('“使”字/使馆格式不支持')
        if t[:1]!=gt[:1]: note.append('首字符错误')
    has_gt=gt is not None and not x
    cat=category(gt) if has_gt else infer_cat(t)+('' if has_gt else '(推断)' if t else '')
    rows.append([i,img,cat,gt or '',t if b else ('(读取失败)' if r['read_fail'] else '(未检出)'),
        ('是' if t==gt else '否') if has_gt else '-', round(max(0,1-ed(t,gt)/len(gt)),3) if has_gt else '-',
        gt_color(gt) if has_gt else '', b['color'] if b else '', round(b['conf'],3) if b else '',
        ('通过' if b.get('strict') else '不通过') if b else '', '；'.join(note),
        round(b['score'],3) if b else '', cb['text'] if cb else '', cb['color'] if cb else ''])
hdr=['序号','图片名称','类别','真值','识别结果','是否正确','字符准确率','颜色真值','识别颜色','识别置信度','格式校验','备注','检测分(模式b)','模式a-直接识别结果','模式a-颜色']
with open('results/results.csv','w',encoding='utf-8-sig',newline='') as f:
    w=csv.writer(f); w.writerow(hdr); w.writerows(rows)
# ---- stats on GT sample ----
def stats(R,filt=None):
    S=collections.defaultdict(lambda:collections.Counter())
    for img,(gt,x) in G.items():
        if x: continue
        b=R[img]['best']; t=b['text'] if b else ''
        for c in (category(gt),'合计'):
            s=S[c]; s['n']+=1; s['det']+=b is not None; s['exact']+=t==gt; s['char']+=max(0,1-ed(t,gt)/len(gt))
            if not gt[0].isdigit(): s['pn']+=1; s['prov']+=t[:1]==gt[:1]
            s['color']+=(b['color'] if b else '')==gt_color(gt); s['strict']+=bool(b and b.get('strict'))
            s['fmt_reject_correct']+= (t==gt and not (b and b.get('strict')))
    return S
SP,SC,S2=stats(P),stats(C),stats(P2)
wb=openpyxl.Workbook(); ws=wb.active; ws.title='明细'; ws.append(hdr)
for r in rows: ws.append(r)
for c in ws[1]: c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='DDEBF7')
red=PatternFill('solid',fgColor='F8CBAD')
for row in ws.iter_rows(min_row=2):
    if row[5].value=='否': row[5].fill=red
for col,wd in zip('ABCDEFGHIJKLMNO',[6,24,18,12,14,8,10,8,8,10,10,46,10,16,10]): ws.column_dimensions[col].width=wd
ws.freeze_panes='A2'
ws2=wb.create_sheet('按类别统计')
h2=['类别','样本数(有真值)','检出数','整牌正确','整牌准确率','字符准确率','省份准确率','颜色准确率','格式校验通过','识别正确但被格式过滤','模式a整牌正确','PAD=2 整牌正确']
ws2.append(h2)
cats=['警用(白)','使馆(黑,使)','新能源小型(绿)','合计']
md=[]
for c in cats:
    s=SP[c]; n=s['n']
    if not n: continue
    r=[c,n,s['det'],s['exact'],f"{s['exact']/n:.1%}",f"{s['char']/n:.1%}",(f"{s['prov']/s['pn']:.1%}" if s['pn'] else '不适用(无省份)'),f"{s['color']/n:.1%}",s['strict'],s['fmt_reject_correct'],SC[c]['exact'],S2[c]['exact']]
    ws2.append(r); md.append(r)
ws2.append([]); ws2.append(['说明：真值为人工抽样标注（随机 120 张，7 张遮挡/模糊不计，共 113 张）；数据集文件名为时间戳 ID，不含真值。主模式 = 模式b（PAD=3 画布 + 检测 + 拉正 + 识别 + 颜色，与板端相同 C++ 代码）。'])
for c in ws2[1]: c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='DDEBF7')
for col in 'ABCDEFGHIJKL': ws2.column_dimensions[col].width=16
# ---- overall ----
ws3=wb.create_sheet('总体统计')
tot=len(imgs); rf=sum(1 for i in imgs if P[i]['read_fail']); det=sum(1 for i in imgs if P[i]['best'])
pred=[P[i]['best']['text'] for i in imgs if P[i]['best']]
cnt=collections.Counter(infer_cat(t) for t in pred)
colc=collections.Counter(P[i]['best']['color'] for i in imgs if P[i]['best'])
strict=sum(1 for i in imgs if P[i]['best'] and P[i]['best'].get('strict'))
s=SP['合计']; n=s['n']
items=[('数据集图片数',tot),('无法读取(损坏)',rf),('模式b 检出车牌',f'{det} ({det/(tot-rf):.1%})'),
 ('格式校验通过(全部)',f'{strict} ({strict/max(det,1):.1%} of 检出)'),
 ('— 有真值抽样 —',''),('抽样数(计入)',n),('整牌准确率(模式b)',f"{s['exact']/n:.1%}"),('字符准确率(模式b)',f"{s['char']/n:.1%}"),
 ('省份准确率(有省份的牌)',f"{s['prov']}/{s['pn']} = {s['prov']/s['pn']:.1%}"),('颜色准确率',f"{s['color']/n:.1%}"),('检出率',f"{s['det']/n:.1%}"),
 ('整牌准确率(模式a 直接识别)',f"{SC['合计']['exact']/n:.1%}"),('字符准确率(模式a)',f"{SC['合计']['char']/n:.1%}"),
 ('— 全量预测分布(按识别结果推断类别) —','')]+[(k,v) for k,v in cnt.most_common()]+[('— 全量识别颜色分布 —','')]+[(k,v) for k,v in colc.most_common()]
ws3.append(['指标','值'])
for k,v in items: ws3.append([k,v])
for c in ws3[1]: c.font=Font(bold=True)
ws3.column_dimensions['A'].width=36; ws3.column_dimensions['B'].width=26
wb.save('results/results.xlsx')
json.dump({'md':md,'items':items},open('results/summary.json','w'),ensure_ascii=False,indent=1)
for r in md: print(r)
for k,v in items: print(k,v)
